#!/usr/bin/env python3
"""Turn a source's downloaded tables into `Community/sources/<id>.json`.

Every source the catalog draws values from is kept the same way: its tables
are downloaded (they stay out of the repo; see the source's
`Community/sources/<id>.yaml` for where from), and this script writes one JSON with
every row of the source — its code, its names as published, and the 16
nutrient fields per 100 g the app knows — one row per line, so a new release
reads as a diff. The JSON is checked in; the compiler reads it, so an
ingredient can name any row of any source by its code and nobody copies a
number by hand. Run this only for a new release of a source.

A value the source leaves blank stays blank: a missing value and a true zero
are different things.

Usage:
    python3 Scripts/sources/extract.py bls <BLS_4_0_Daten_2025_DE.xlsx>
    python3 Scripts/sources/extract.py ciqual-2020 <Table Ciqual 2020_ENG_2020 07 07.xls>
    python3 Scripts/sources/extract.py ciqual-2025 <Table Ciqual 2025_FR_….xlsx> <alim_2025_….xml>
    python3 Scripts/sources/extract.py usda-sr-legacy <FoodData_Central_sr_legacy_food_csv_2018-04.zip>
    python3 Scripts/sources/extract.py usda-foundation-foods <FoodData_Central_foundation_food_csv_2021-10-28.zip>
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SOURCES = REPO_ROOT / "Community" / "sources"

# The app's 16 fields, in the order every CSV writes them.
FIELDS = [
    "kcal", "proteinG", "fatG", "saturatedFatG", "carbsG", "sugarG", "fiberG",
    "sodiumMg", "vitaminAMcg", "vitaminCMg", "vitaminDMcg", "vitaminEMg",
    "calciumMg", "ironMg", "magnesiumMg", "potassiumMg",
]


def number(value) -> str:
    """A value as the CSV writes it: the source's own precision, no float noise."""
    if value is None or value == "":
        return ""
    rounded = round(float(value), 4)
    return str(int(rounded)) if rounded == int(rounded) else repr(rounded)


def row_json(row: dict) -> dict:
    """A row as `Community/sources/<id>.json` keeps it: the names, and only the
    values the source states."""
    out = {"name": row["name"]}
    if row.get("nameEnglish") and row["nameEnglish"] != row["name"]:
        out["nameEnglish"] = row["nameEnglish"]
    out["per100g"] = {field: float(row[field]) if "." in row[field] else int(row[field])
                      for field in FIELDS if row.get(field, "") != ""}
    return out


def write(source_id: str, rows: list[dict]) -> Path:
    path = SOURCES / f"{source_id}.json"
    write_rows(path, {row["code"]: row_json(row) for row in rows})
    return path


def write_rows(path: Path, rows: dict[str, dict]) -> None:
    """One row per line, keyed by the source's code, in code order."""
    lines = [f" {json.dumps(code, ensure_ascii=False)}: {json.dumps(row, ensure_ascii=False)}"
             for code, row in sorted(rows.items())]
    path.write_text('{"rows": {\n' + ",\n".join(lines) + "\n}}\n", encoding="utf-8")


# --------------------------------------------------------------------------
# BLS 4.0: one workbook, a header row naming each nutrient by its BLS code.

BLS_SHEET = "BLS_4_0_Daten_2025_DE"

# Field -> BLS nutrient code. VITAA falls back to VITA where it is blank: the
# BLS fills only one of its two vitamin A measures for a lot of rows.
BLS_NUTRIENTS = {
    "kcal": "ENERCC", "proteinG": "PROT625", "fatG": "FAT", "saturatedFatG": "FASAT",
    "carbsG": "CHO", "sugarG": "SUGAR", "fiberG": "FIBT", "sodiumMg": "NA",
    "vitaminAMcg": "VITAA", "vitaminCMg": "VITC", "vitaminDMcg": "VITD", "vitaminEMg": "VITE",
    "calciumMg": "CA", "ironMg": "FE", "magnesiumMg": "MG", "potassiumMg": "K",
}
BLS_VITAMIN_A_FALLBACK = "VITA"


def bls_columns(header: tuple, codes: set[str]) -> dict[str, int]:
    """Each nutrient code's value column, found by its header — "ENERCC
    Energie (Kilokalorien) [kcal/100g]" — not by position: the workbook is
    "value, Datenherkunft, Referenz" triplets, and which nutrients a release
    exports shifts every offset."""
    found = {}
    for index, cell in enumerate(header):
        text = str(cell or "")
        for code in codes - found.keys():
            if text.startswith(code + " "):
                found[code] = index
    missing = codes - found.keys()
    if missing:
        raise SystemExit(f"No column for the BLS code(s) {sorted(missing)} in the header")
    return found


def bls(xlsx: Path) -> list[dict]:
    import openpyxl
    codes = set(BLS_NUTRIENTS.values()) | {BLS_VITAMIN_A_FALLBACK}
    sheet = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)[BLS_SHEET]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    columns = bls_columns(header, codes)
    name_column = header.index("Lebensmittelbezeichnung")
    english_column = header.index("Food name")
    out = []
    for row in rows:
        code = row[0]
        if not code:
            continue
        values = {}
        for field, nutrient in BLS_NUTRIENTS.items():
            value = row[columns[nutrient]]
            if value in (None, "") and nutrient == "VITAA":
                value = row[columns[BLS_VITAMIN_A_FALLBACK]]
            # A blank stays blank; a stray text ("Spur") is no number either.
            if isinstance(value, (int, float)):
                values[field] = number(value)
        out.append({
            "code": str(code).strip(),
            "name": str(row[name_column] or "").strip(),
            "nameEnglish": str(row[english_column] or "").strip(),
            **{field: values.get(field, "") for field in FIELDS},
        })
    return out


# --------------------------------------------------------------------------
# Ciqual (Anses): one sheet, a header per constituent with its unit. A value
# is a number with a decimal comma, "-" for unknown, "traces" or "< x" below
# the limit of quantification — neither of the last two is a zero, so both
# stay blank (Community/README.md, rule 5).

def ciqual_value(cell) -> str:
    if isinstance(cell, (int, float)):
        return number(cell)
    text = str(cell or "").strip()
    if not text or text == "-" or text.startswith("<") or text.lower() == "traces":
        return ""
    return number(float(text.replace(",", ".")))


def ciqual_header(cell) -> str:
    return " ".join(str(cell or "").split()).lower()


# Field -> the header a column starts with, English (2020) or French (2025).
# Ciqual 2020 has no vitamin A activity, only retinol and beta-carotene apart;
# the rows taken from it never carried vitamin A, and summing the two would
# need a conversion factor the table does not state. 2025 states the activity.
CIQUAL_COLUMNS = {
    "kcal": ["energy, regulation eu no 1169/2011 (kcal", "energie, règlement ue n° 1169 2011 (kcal"],
    "proteinG": ["protein (g", "protéines, n x facteur de jones"],
    "fatG": ["fat (g", "lipides"],
    "saturatedFatG": ["fa saturated", "ag saturés"],
    "carbsG": ["carbohydrate (g", "glucides"],
    "sugarG": ["sugars (g", "sucres"],
    "fiberG": ["fibres (g", "fibres alimentaires"],
    "sodiumMg": ["sodium (mg"],
    "vitaminAMcg": ["activité vitaminique a"],
    "vitaminCMg": ["vitamin c (mg", "vitamine c"],
    "vitaminDMcg": ["vitamin d (µg", "vitamine d (µg"],
    # Vitamin E as a whole, as 2020 states it — not 2025's alpha-tocopherol
    # alone, which is a different (and often the only filled) column.
    "vitaminEMg": ["vitamin e (mg", "vitamine e (mg"],
    "calciumMg": ["calcium (mg"],
    "ironMg": ["iron (mg", "fer (mg"],
    "magnesiumMg": ["magnesium (mg", "magnésium (mg"],
    "potassiumMg": ["potassium (mg"],
}


def ciqual_rows(header: list, rows, english_names: dict[str, str] | None = None) -> list[dict]:
    names = [ciqual_header(cell) for cell in header]
    columns = {}
    for field, prefixes in CIQUAL_COLUMNS.items():
        for index, name in enumerate(names):
            if any(name.startswith(prefix) for prefix in prefixes):
                columns[field] = index
                break
    code_column = names.index("alim_code")
    name_column = next(i for i, n in enumerate(names) if n in ("alim_nom_eng", "alim_nom_fr"))
    out = []
    for row in rows:
        raw_code = row[code_column]
        if raw_code in (None, ""):
            continue
        code = str(int(float(raw_code)))
        own_name = str(row[name_column] or "").strip()
        out.append({
            "code": code,
            "name": own_name,
            "nameEnglish": (english_names or {}).get(code, own_name if names[name_column] == "alim_nom_eng" else ""),
            **{field: ciqual_value(row[index]) if field in columns and (index := columns[field]) is not None else ""
               for field in FIELDS},
        })
    return out


def ciqual_2020(xls: Path) -> list[dict]:
    """The English workbook, whose names the rows cite."""
    import xlrd
    sheet = xlrd.open_workbook(xls).sheet_by_index(0)
    return ciqual_rows(sheet.row_values(0), (sheet.row_values(r) for r in range(1, sheet.nrows)))


def ciqual_2025(xlsx: Path, alim_xml: Path) -> list[dict]:
    """The French workbook (the only one there is) and `alim_*.xml` for the
    English names, which the rows cite."""
    import openpyxl
    import xml.etree.ElementTree as ET
    english = {}
    for food in ET.parse(alim_xml).getroot().iter("ALIM"):
        code = (food.findtext("alim_code") or "").strip()
        if code:
            english[code] = (food.findtext("alim_nom_eng") or "").strip()
    sheet = openpyxl.load_workbook(xlsx, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    return ciqual_rows(header, rows, english)


# --------------------------------------------------------------------------
# USDA FoodData Central: the CSV download, as a zip. `food.csv` names the
# foods by FDC ID, `food_nutrient.csv` holds one amount per food and nutrient
# (a nutrient with no row is unknown), `nutrient.csv` names the nutrients.

# Field -> nutrient ids in order of preference. Energy: the kcal the release
# states; Foundation Foods states it only through the Atwater factors. Fibre:
# Foundation Foods measures it by AOAC 2011.25 where SR Legacy has the total.
#
# No carbohydrate: USDA states it "by difference", which counts the fibre in,
# while the app's carbohydrate — the BLS's and Ciqual's — leaves it out. The
# rows taken by hand never carried it either, and an empty field is honest
# where a converted one would only look complete.
USDA_NUTRIENTS = {
    "kcal": [1008, 2047, 2048],
    "proteinG": [1003],
    "fatG": [1004],
    "saturatedFatG": [1258],
    "sugarG": [2000, 1063],
    "fiberG": [1079, 2033],
    "sodiumMg": [1093],
    "vitaminAMcg": [1106],
    "vitaminCMg": [1162],
    "vitaminDMcg": [1114],
    "vitaminEMg": [1109],
    "calciumMg": [1087],
    "ironMg": [1089],
    "magnesiumMg": [1090],
    "potassiumMg": [1092],
}


def usda(zip_path: Path) -> list[dict]:
    import io
    import zipfile
    wanted = {nid for ids in USDA_NUTRIENTS.values() for nid in ids}
    with zipfile.ZipFile(zip_path) as archive:
        def table(name):
            member = next(m for m in archive.namelist() if m.endswith("/" + name) or m == name)
            return csv.DictReader(io.TextIOWrapper(archive.open(member), encoding="utf-8"))
        # Foundation Foods ships its samples and sub-samples as foods of their
        # own; only the food itself is a row anybody would name.
        foods = {row["fdc_id"]: row["description"].strip() for row in table("food.csv")
                 if row.get("data_type") in ("sr_legacy_food", "foundation_food")}
        amounts: dict[str, dict[int, str]] = {}
        for row in table("food_nutrient.csv"):
            nid = int(row["nutrient_id"])
            if nid in wanted and row["fdc_id"] in foods and row["amount"] != "":
                amounts.setdefault(row["fdc_id"], {})[nid] = row["amount"]
    out = []
    for fdc_id, description in foods.items():
        values = amounts.get(fdc_id, {})
        row = {"code": fdc_id, "name": description, "nameEnglish": description}
        for field, ids in USDA_NUTRIENTS.items():
            row[field] = next((number(values[nid]) for nid in ids if nid in values), "")
        out.append(row)
    return out


EXTRACTORS = {
    "bls": bls,
    "ciqual-2020": ciqual_2020,
    "ciqual-2025": ciqual_2025,
    "usda-sr-legacy": usda,
    "usda-foundation-foods": usda,
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", choices=sorted(EXTRACTORS))
    parser.add_argument("paths", nargs="+", type=Path, help="the downloaded file(s)")
    args = parser.parse_args()
    rows = EXTRACTORS[args.source](*args.paths)
    path = write(args.source, rows)
    print(f"Wrote {len(rows)} rows to {path.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
